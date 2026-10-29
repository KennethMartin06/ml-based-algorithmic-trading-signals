"""Build the speaker-notes handbook (presentation/speaker_notes.docx and .pdf).

One page per slide: a thumbnail, what every element on the slide means, the script to say (the deck's own
speaker notes), the numbers to remember and likely examiner questions with answers. Appendices: a glossary,
a click-by-click live-demo script for the dashboard and a one-page numbers cheat sheet.

Numbers come from the same `Facts` the deck uses, so the handbook and the slides cannot disagree.
Needs the slide renders from render_preview.py (presentation/preview/slide-NN.png).

Run from the project root:  python presentation/scripts/build_speaker_notes.py
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_deck import LIVE, PRES, SHORT, TICKERS, Facts, build  # noqa: E402

PREVIEW = PRES / "preview"
OUT = PRES / "speaker_notes.docx"
WPM = 130  # speaking pace used for the timings
INK, AMBER, MUTED = RGBColor(0x1B, 0x22, 0x2C), RGBColor(0xB3, 0x6B, 0x00), RGBColor(0x5B, 0x67, 0x75)


def strategy_names(F: Facts) -> tuple[list[str], list[str]]:
    t = F.trade
    bh = t[t.strategy == "Buy & Hold"].set_index("ticker").sharpe
    conf = t[t.strategy.str.endswith("(conf.)")].copy()
    conf["bh"] = conf.ticker.map(bh)
    lab = lambda r: f"{r.ticker} {SHORT[r.strategy.replace(' (conf.)', '')]}"  # noqa: E731
    higher = [lab(r) for r in conf.itertuples() if r.sharpe > r.bh + 1e-9]
    same = [lab(r) for r in conf.itertuples() if abs(r.sharpe - r.bh) <= 1e-9]
    return higher, same


def content(F: Facts) -> list[dict]:
    """Per-slide explanation, numbers and Q&A. Order matches the deck."""
    sp = F.split
    est, lo, hi = F.bh_ci_spy
    higher, same = strategy_names(F)
    rel = F.acc_reliable.iloc[0]
    worst = F.sharpe_reliable.iloc[0]
    P = F.params
    pc = F.pers_cost
    sh = F.spy_test_share
    wf = F.wf_table
    am = ", ".join(f"{tk} {SHORT[s.replace(' (argmax)', '')]} {a:.2f} vs {b:.2f}" for tk, s, a, b in F.argmax_higher)
    lo_r, hi_r = F.vol_red
    new_days = LIVE["test_days"] - sp["test"]["rows"]
    return [
        dict(  # 1
            shows=["The project title, subtitle, the two presenters and the two mentors.",
                   "The amber line is real data: SPY's daily adjusted closing price from 2010 to 2025, the same "
                   "series the models were trained and tested on. It draws itself left to right when the slide "
                   "opens and comes back on the last slide."],
            numbers=["Data: Yahoo Finance daily prices, 2010-01-01 to 2025-12-31",
                     "Instruments: SPY (S&P 500 ETF), Apple (AAPL), Microsoft (MSFT)"],
            qa=[("Why SPY, Apple and Microsoft?",
                 "They are very liquid, have long, clean daily histories that are free to download, and SPY "
                 "represents the whole US market. The honest caveat: we picked them knowing they did well, which is "
                 "survivorship bias (slide 20).")]),
        dict(  # 2
            shows=["Left card: the classic hand-written rule (moving-average crossover). Buy when the 10-day "
                   "average is above the 50-day average, otherwise stay in cash.",
                   f"Right card: the learned rule. A model looks at {F.n_features} features and estimates "
                   "P(Up), the probability that the next move is up. We buy only when P(Up) is at least a "
                   "threshold τ (tau).",
                   "Bottom line: both approaches get identical data, execution timing and costs, so any difference "
                   "comes from the signal. The grey line names the three instruments.",
                   "The amber line from slide 1 morphs down to the bottom of this slide."],
            numbers=[f"τ is chosen from {F.grid_n} values between {F.grid_min:.2f} and {F.grid_max:.2f}",
                     f"Training data 2011–2019, threshold chosen on 2020–2021"],
            qa=[("What exactly is τ?",
                 "A probability cut-off. If the model says P(Up) ≥ τ we are invested the next day, otherwise in "
                 "cash. It is chosen on the 2020–21 validation period by Sharpe ratio, never on the test period."),
                ("Why cash instead of short-selling?",
                 "Long-only is the default because it is realistic for a retail investor and avoids borrow costs. "
                 "The dashboard can also show long/short.")]),
        dict(  # 3
            shows=["Four cards, each a common way a trading backtest gives falsely good results. The small grey "
                   "line at the bottom of each card is the symptom you would see.",
                   "01 Look-ahead bias: using information that was not available at decision time.",
                   "02 Leaky validation: random K-fold cross-validation on time series trains on the future.",
                   "03 Free trading: ignoring fees, slippage and when orders really fill.",
                   f"04 Weak baselines: a model at {F.acc_max * 100:.0f}% accuracy sounds good, but 'always predict "
                   f"Up' already scores {F.maj_spy * 100:.1f}% on SPY."],
            numbers=[f"'Always predict Up' accuracy on SPY test: {F.maj_spy * 100:.1f}%"],
            qa=[("Give a concrete example of look-ahead bias.",
                 "Normalising a feature with the mean of the whole dataset (which includes future years), or "
                 "deciding with today's close and assuming you also traded at today's close."),
                ("Why is K-fold wrong for time series?",
                 "Random folds put days from the future into the training set of earlier test days, so the model "
                 "learns about regimes it should not know yet.")]),
        dict(  # 4
            shows=["Chart: a deliberately naive strategy, 'persistence' (tomorrow repeats today). The x-axis is "
                   "the trading cost in basis points, drawn to scale from 0 to 40. The y-axis is the Sharpe ratio. "
                   "Cyan = SPY, violet = AAPL, teal = MSFT (AAPL and MSFT nearly overlap). The dashed white line is "
                   "buy-and-hold on SPY, flat because it trades only once.",
                   f"Top-right number {F.moc_from:.2f} → {F.moc_to:.2f}: the same Random Forest on AAPL, with only "
                   "the fill price changed from next open to the same day's close (a sensitivity run).",
                   f"Second number {lo:.2f} … {hi:.2f}: the 95% bootstrap interval for buy-and-hold's Sharpe "
                   f"ratio on SPY ({est:.2f}). Four years of data is a small sample.",
                   "Footnote: definitions of the Sharpe ratio and of a basis point."],
            numbers=[f"Persistence Sharpe at 0 bps: SPY {pc['SPY'][0]:.2f}, AAPL {pc['AAPL'][0]:.2f}, "
                     f"MSFT {pc['MSFT'][0]:.2f}",
                     f"At 40 bps: SPY {pc['SPY'][-1]:.2f}, AAPL {pc['AAPL'][-1]:.2f}, MSFT {pc['MSFT'][-1]:.2f}",
                     "1 bp (basis point) = 0.01%"],
            qa=[("What is a Sharpe ratio?",
                 "Annualised average return divided by annualised volatility, roughly return per unit of risk. "
                 "We use a risk-free rate of 0, so it is simply return divided by volatility."),
                ("Why does the fill price matter so much?",
                 "With a same-day-close fill you trade at the price you used to compute the signal, which is not "
                 "really achievable. Next-open fills are realistic, and overnight gaps change the return you get.")]),
        dict(  # 5
            shows=["The four objectives copied from our synopsis, each with a status chip. Objectives 1–3 are "
                   "delivered. Objective 4 (news sentiment via retrieval-augmented generation, RAG) is future work.",
                   "Bottom line: what we added beyond the synopsis."],
            numbers=["Extra models: Logistic Regression, Random Forest, soft-vote ensemble",
                     "Extra evaluation: walk-forward, bootstrap confidence intervals, live dashboard"],
            qa=[("Why didn't you build the RAG part?",
                 "Doing it correctly needs a time-stamped news archive for 2010–2025, which is not reliably "
                 "available for free. News published after the decision time would leak the future and invalidate "
                 "the comparison. We designed it instead (slide 21).")]),
        dict(  # 6
            shows=["Left column: the four traps from slide 3. Arrows lead to our countermeasure for each.",
                   "Point-in-time features; chronological splits with a 2-day purge; next-open execution with "
                   f"{F.cost_total} bps costs; four baselines plus bootstrap confidence intervals.",
                   "Footer: the main results use all four safeguards. Sensitivity runs relax one at a time and are "
                   "labelled."],
            numbers=[f"Costs: {F.fee_bps} bps fee + {F.slip_bps} bps slippage = {F.cost_total} bps per unit "
                     "of turnover",
                     "Purge: 2 rows before each boundary"],
            qa=[("What does 'purge' mean?",
                 "Each label uses the open two days ahead, so the last two rows before a boundary have labels that "
                 "reach into the next segment. Those rows are dropped."),
                ("Where does 15 bps come from?",
                 "It is an illustrative retail cost (10 fee + 5 slippage). Slide 19 shows results from 0 to 40 bps.")]),
        dict(  # 7
            shows=["Each box is one stage of the pipeline, with the Python file it lives in.",
                   "Row 1 (appears automatically): data, features, targets, split.",
                   "Click 1 shows row 2: models, signal rule, backtester, evaluation.",
                   "Click 2 shows the outputs band: tables, figures, reports, dashboard."],
            numbers=[f"{F.n_features} features · 5 model families · {F.n_figures} generated figures"],
            qa=[("How long does it take to run?",
                 f"About {F.minutes} minutes on a laptop CPU for all three tickers, including walk-forward and "
                 "robustness runs."),
                ("Is anything typed by hand?",
                 "No. Every table, figure and report, including this deck, is generated from the saved results.")]),
        dict(  # 8
            shows=["Four columns: data, modelling, evaluation and delivery tools, with the exact versions recorded "
                   "for the reported run."],
            numbers=[f"Python {F.env['python']}, pandas {F.env['pandas']}, scikit-learn {F.env['scikit_learn']}, "
                     f"XGBoost {F.env['xgboost']}, PyTorch {F.env['torch'].replace('+cpu', '')}",
                     f"{F.n_tests} pytest tests · fixed seed {F.seed}"],
            qa=[("Why XGBoost and PyTorch?",
                 "XGBoost is the strongest standard tool for tabular data. PyTorch implements the LSTM sequence "
                 "model that the synopsis proposed.")]),
        dict(  # 9
            shows=["A time axis from 2011 to 2026. The three bars wipe in left to right: train (teal), validation "
                   "(amber), test (white). 'Purge' markers sit at the boundaries.",
                   "The cards below say what each period is used for."],
            numbers=[f"Train {sp['train']['start']} → {sp['train']['end']}: {sp['train']['rows']:,} days",
                     f"Validation {sp['val']['start']} → {sp['val']['end']}: {sp['val']['rows']} days",
                     f"Test {sp['test']['start']} → {sp['test']['end']}: {sp['test']['rows']:,} days"],
            qa=[("Why does training start in 2011 if the data starts in 2010?",
                 "The longest features (200- and 252-day windows) need about a year of history before their first "
                 "complete value."),
                ("Wasn't 2020–21 an unusual validation period?",
                 "Yes. It was a strong bull market, so τ was often chosen low enough to stay invested every day. "
                 "We report that openly on slides 18 and 20.")]),
        dict(  # 10
            shows=["The three bars from slide 9 morph into the last of ten rows. Each row is one walk-forward fold, "
                   "labelled with its test year: training data (expanding), the validation year just before it, and "
                   "the test year traded out-of-sample."],
            numbers=[f"{len(F.wf_years)} folds ({F.wf_years[0]}–{F.wf_years[-1]}) × 3 tickers",
                     "Train up to Y−2, validate on Y−1, trade Y"],
            qa=[("How is this different from the hold-out test?",
                 "It is a separate check: ten one-year out-of-sample periods instead of one four-year period. The "
                 "tabular models are re-tuned in every fold (smaller grids); the LSTM keeps its architecture and "
                 "only re-tunes class weighting.")]),
        dict(  # 11
            shows=["Left: the 42 features split into 5 groups (bar segment widths are the counts), with examples.",
                   "Right: the timing of one decision. Features at close t, enter at open t+1, exit or re-decide at "
                   "open t+2.",
                   "The equation is the return we predict. The three chips are the classes (±0.1% neutral band).",
                   "Bottom line: how often each class occurred on SPY in the test period."],
            numbers=[f"Groups: returns {F.fgroups['returns']}, technical {F.fgroups['technical']}, volatility "
                     f"{F.fgroups['volatility']}, volume {F.fgroups['volume']}, regime {F.fgroups['regime']}",
                     f"SPY test days: Up {sh['Up'] * 100:.1f}%, Down {sh['Down'] * 100:.1f}%, Neutral "
                     f"{sh['Neutral'] * 100:.1f}%"],
            qa=[("Why three classes and not just up/down?",
                 f"The ±{F.neutral * 100:.1f}% neutral band removes tiny moves that are mostly noise. A binary "
                 "up/down variant was also tested (slide 19)."),
                ("Why only ratios?",
                 "Adjusted prices are rewritten backwards after splits and dividends, so raw price levels can carry "
                 "future information. Ratios are unaffected, and a test checks that features never change when "
                 "future rows are removed.")]),
        dict(  # 12
            shows=["Five cards, one per model family, with the hyperparameters that tuning selected. The amber "
                   "card (XGBoost) is one of the two core models named in the synopsis.",
                   "Values that differ by ticker are written SPY / AAPL / MSFT. The ensemble lists its two members "
                   "per ticker.",
                   "Bottom: the main insight, and the four baselines."],
            numbers=[f"LR C = {P['SPY']['logistic_regression']['C']:g} (strongest regularisation in the grid)",
                     f"RF depth {P['SPY']['random_forest']['max_depth']}, 300 trees; XGB depth "
                     f"{P['SPY']['xgboost']['max_depth']}, η {P['SPY']['xgboost']['learning_rate']:g}, "
                     f"{P['SPY']['xgboost']['n_estimators']} trees",
                     "LSTM: 1 layer, hidden " + "/".join(str(P[t]["lstm"]["hidden_size"]) for t in TICKERS)
                     + ", 10-day window"],
            qa=[("Why tune on log loss instead of accuracy?",
                 "Trades are driven by probabilities. Log loss rewards well-calibrated probabilities, while "
                 "accuracy only checks the most likely class."),
                ("Why such a small LSTM?",
                 "There are only about 2,200 training days. Bigger networks memorise noise, so we use one small "
                 "layer with early stopping on a purged slice of the training period.")]),
        dict(  # 13
            shows=["Left: a probability bar from 0 to 1. Below τ = cash, above τ = long. The marker shows SPY "
                   "XGBoost's τ.",
                   "Right: the two backtest equations applied every day. Costs are charged on the change in "
                   "position, then the day's return is applied.",
                   "Chips: every metric reported for every strategy."],
            numbers=[f"SPY XGBoost τ = {F.tau_spy_xgb}",
                     f"Chosen τ across all 15 model–ticker pairs: {F.thr_min:.2f}–{F.thr_max:.2f}",
                     f"Starting capital {F.capital:,}"],
            qa=[("Why decide at the close but fill at the next open?",
                 "You only know the close once the market has closed. The earliest realistic trade is the next "
                 "morning's open."),
                ("What is the argmax rule?",
                 "A simpler rule without τ: invest whenever Up is the most likely of the three classes.")]),
        dict(  # 14
            shows=["Six tiles of facts about the engineering: tests, one-command reproducibility, runtime, "
                   "generated figures, dashboard and bootstrap resamples."],
            numbers=[f"{F.n_tests} tests · {F.minutes} min full run · {F.n_figures} figures · "
                     f"{F.n_boot:,} bootstrap resamples (block {F.block})",
                     "Rebuild command: python -m scripts.run_experiment"],
            qa=[("What does a truncation-invariance test do?",
                 "It computes the features on the full history and again on history cut at a date t. The values "
                 "at t must be identical; if not, a feature used future data."),
                ("Is the run deterministic?",
                 f"Yes on the same machine and library versions (seed {F.seed}). Other versions can change late "
                 "decimals.")]),
        dict(  # 15
            shows=["Three real screenshots of our Streamlit dashboard, captured 24 Oct 2026.",
                   "Large: the interactive backtest (pick a model, move τ, change costs or sizing; numbers recompute "
                   "live).",
                   "Right top: the walk-forward chart. The dashed line (buy-and-hold) finishes above all five ML "
                   "strategies.",
                   "Right bottom: the metric cards from a live training run, which the next slide opens.",
                   "Good moment for the live demo: see Appendix B."],
            numbers=["Six tabs: Overview, Interactive backtest, Predictions, Walk-forward & robustness, Train a "
                     "model live, Figures & docs"],
            qa=[("Are these numbers recomputed or pre-saved?",
                 "The backtest is recomputed live by our own backtester from saved out-of-sample probabilities. "
                 "The 'Train a model live' tab downloads fresh data and retrains.")]),
        dict(  # 16
            shows=["The live screenshot grows from the thumbnail into full size (Morph).",
                   "Right: three facts from that one live run on 24 Oct 2026, and the takeaway.",
                   "The header says clearly that this is one run, not part of the reported results."],
            numbers=[f"Pipeline time {LIVE['seconds']} s",
                     f"Test window to {LIVE['test_end_label']}: {LIVE['test_days']:,} days ({new_days} more than "
                     "the study's test set)",
                     f"Sharpe {LIVE['sharpe']:.2f} vs buy-and-hold {LIVE['bh_sharpe']:.2f}; accuracy "
                     f"{LIVE['acc_gap_pp']} points below the majority class"],
            qa=[("If you press the button today, will the numbers match the slide?",
                 "Not exactly. The data updates every trading day, so the test window grows and the numbers move "
                 "slightly. The slide records the run of 24 Oct 2026."),
                ("Is the live window 'unseen data'?",
                 f"Only partly. {sp['test']['rows']:,} of its {LIVE['test_days']:,} test days are the study's own "
                 "test period, which is why the title says 'ten more months of data' and not 'unseen data'.")]),
        dict(  # 17
            shows=["Each bar = a model's test accuracy minus the accuracy of 'always predict Up' (the majority "
                   "class), in percentage points. Colours are tickers. The thin whiskers are 95% bootstrap "
                   "intervals.",
                   "Almost every whisker crosses zero: no reliable difference.",
                   "Right: the accuracy range of our models, the baseline's range, and the one reliable difference "
                   "(which is worse)."],
            numbers=[f"Model accuracy {F.acc_min * 100:.1f}–{F.acc_max * 100:.1f}% vs baseline "
                     f"{F.maj_min * 100:.1f}–{F.maj_max * 100:.1f}%",
                     f"Only reliable difference: {rel.ticker} {rel.strategy}, {abs(rel.estimate) * 100:.1f} pp worse"],
            qa=[("Isn't 51% better than 50%?",
                 "Only by about 1 point, and the 95% interval for that difference includes zero. It is "
                 "indistinguishable from the baseline."),
                ("What is a bootstrap confidence interval?",
                 f"We resample the test days in blocks of {F.block} days (to keep their time structure) "
                 f"{F.n_boot:,} times and recompute the metric each time. The middle 95% of those values is the "
                 "interval.")]),
        dict(  # 18
            shows=["Each bar = one τ-rule strategy's Sharpe ratio minus buy-and-hold's, grouped by ticker. Whiskers "
                   "are 95% bootstrap intervals. Bars to the right are better than buy-and-hold.",
                   "'= buy-and-hold' marks strategies that were invested every single day, so they are identical to "
                   "it.",
                   "The coral bar is the only reliable difference, and it is worse.",
                   "The note under the chart discloses the argmax rule's results."],
            numbers=[f"Higher Sharpe than buy-and-hold ({F.n_higher}/15, none reliable): {', '.join(higher)}",
                     f"Identical to buy-and-hold ({F.n_identical}/15): {', '.join(same)}",
                     f"Reliably worse: {worst.ticker} {worst.model}",
                     f"Naive persistence on SPY: {F.pers_spy_cum * 100:.0f}% after {F.pers_spy_trades} trades",
                     f"Argmax rule: {len(F.argmax_higher)} of {F.n_argmax} above buy-and-hold ({am}), no intervals"],
            qa=[("Why are five strategies exactly equal to buy-and-hold?",
                 "The τ chosen on the bullish 2020–21 validation period was low enough that P(Up) always cleared "
                 "it, so they never left the market."),
                ("Two argmax strategies beat buy-and-hold. Doesn't that contradict the title?",
                 "The title is about the 15 τ-rule strategies we tested with intervals. For argmax we did not "
                 "compute intervals, so we make no claim either way. Gaps of about 0.1 in Sharpe sat well inside "
                 "the intervals of the τ-rule strategies, but we did not test the argmax ones.")]),
        dict(  # 19
            shows=["Chart: walk-forward 2016–2025 Sharpe ratios for buy-and-hold (white), MA crossover (grey) and "
                   "the best ML strategy per ticker chosen with hindsight (amber), the most favourable comparison "
                   "for ML.",
                   "Big number: how many ML strategies beat buy-and-hold over the ten years.",
                   "Cards: robustness checks run on the 2022–2025 hold-out test (costs, sizing, alternative target "
                   "and fill)."],
            numbers=["Walk-forward Sharpe (B&H / MA / best ML): " + "; ".join(
                f"{t} {wf[t][0]:.2f}/{wf[t][1]:.2f}/{wf[t][2]:.2f}" for t in TICKERS),
                     f"ML beating buy-and-hold in walk-forward: {F.wf_beat}/{F.wf_n}",
                     f"Costs: {F.cost_beat_lo[0]}/15 above B&H at 0 bps, {F.cost_beat_hi[0]}/15 at 40 bps",
                     f"Volatility targeting: max drawdown −{lo_r * 100:.0f} to −{hi_r * 100:.0f}% at a similar Sharpe",
                     f"Alternative target or fills: {F.var_lo}–{F.var_hi} of {F.var_n} above B&H"],
            qa=[("Why pick the best ML model with hindsight?",
                 "To give ML every advantage. Even the cherry-picked model loses to buy-and-hold."),
                ("Doesn't volatility targeting just reduce exposure?",
                 "Yes, that is the point. It lowers drawdowns by holding less when markets are volatile, at a "
                 "similar Sharpe and a lower return. It is risk control, not prediction.")]),
        dict(  # 20
            shows=["Left: engineering challenges we handled. Right: limitations we state openly. These are the "
                   "points an examiner is most likely to probe."],
            numbers=["~2,200 training days per ticker · flat 15 bps cost model · no news or macro data"],
            qa=[("What is survivorship bias here?",
                 "SPY, Apple and Microsoft are known winners. Strategies tested on companies that later failed "
                 "could behave differently."),
                ("What is the multiple-testing problem?",
                 "When you try many models × thresholds × assets, some look good by chance. A correction such as "
                 "the deflated Sharpe ratio would adjust for that; we list it as future work.")]),
        dict(  # 21
            shows=["Top: the five-step design for the news-sentiment (RAG) extension, built to avoid look-ahead.",
                   "Bottom: three directions for future work: data, statistics and practice."],
            numbers=["Key rule: only retrieve news published before the close of day t"],
            qa=[("What is RAG?",
                 "Retrieval-augmented generation. For a ticker and date, relevant documents (news) are retrieved "
                 "and fed to a model. Here, a financial sentiment model such as FinBERT scores them, and the scores "
                 "become daily features."),
                ("What is FinBERT?",
                 "A BERT language model fine-tuned on financial text to classify sentiment as positive, negative "
                 "or neutral.")]),
        dict(  # 22
            shows=["The three claims we can defend, each with its evidence underneath."],
            numbers=[f"0 of {F.wf_n} walk-forward strategies beat buy-and-hold · {F.n_tests} tests · drawdown "
                     f"−{lo_r * 100:.0f} to −{hi_r * 100:.0f}% with volatility targeting"],
            qa=[("So is machine learning useless for trading?",
                 "That is not our claim. At a daily horizon, with price-and-volume features on three very liquid "
                 "instruments and after realistic costs, it did not reliably beat simple baselines. Other horizons, "
                 "data (news, order books) or markets may differ; that is the future work."),
                ("Then what is the contribution?",
                 "An evaluation framework that makes a negative result trustworthy: leakage-safe, cost-aware, "
                 "tested and reproducible. Many published positive results do not survive these checks.")]),
        dict(  # 23
            shows=["Thank you, the repository, the command that starts the demo, and the team. The price line "
                   "morphs back to full size."],
            numbers=["Demo: streamlit run app/streamlit_app.py (http://localhost:8501)"],
            qa=[]),
    ]


GLOSSARY = [
    ("Sharpe ratio", "Annualised mean return ÷ annualised volatility (risk-free rate 0 here). Return per unit of "
                     "risk; higher is better."),
    ("Basis point (bp)", "0.01%. 15 bps = 0.15% of the traded amount."),
    ("Turnover", "How much of the portfolio changes hands. Going from cash to fully invested is 1 unit."),
    ("Buy-and-hold", "Buy on the first day and hold to the end. The benchmark to beat."),
    ("MA crossover (10/50)", "Invested when the 10-day average price is above the 50-day average."),
    ("Majority class", "Always predicting the most common class (Up). The accuracy baseline."),
    ("Naive persistence", "Predict that tomorrow moves like today."),
    ("τ (tau)", "Probability threshold. Invest when P(Up) ≥ τ. Chosen on validation."),
    ("Argmax rule", "Invest when Up is the most likely class (no threshold)."),
    ("Look-ahead bias", "Using information that was not available at decision time."),
    ("Purge", "Dropping the last rows before a split boundary whose labels reach into the next period."),
    ("Walk-forward", "Repeatedly retrain on the past and test on the next year, moving forward in time."),
    ("Bootstrap CI", "Resample the test days (in blocks) many times and recompute the metric; the middle 95% is "
                     "the confidence interval."),
    ("Log loss", "Scores predicted probabilities; punishes confident wrong predictions. Used for tuning."),
    ("LSTM", "Long short-term memory network, a recurrent neural network for sequences."),
    ("Soft-vote ensemble", "Average the probabilities of several models."),
    ("Volatility targeting", "Scale the position so the portfolio's volatility stays near a target (15%/yr)."),
    ("Max drawdown", "Largest fall from a previous peak, in %."),
    ("Survivorship bias", "Testing only on assets that survived and did well."),
    ("RAG", "Retrieval-augmented generation: retrieve relevant documents and feed them to a model."),
]

DEMO = [
    ("Before the viva", "Start the dashboard (see the launch line at the end of this appendix) and open "
                        "http://localhost:8501. The first load takes about 10–20 seconds. The live-training tab "
                        "needs internet; everything else runs offline from saved results."),
    ("1 · Overview tab", "Point at the headline cards: buy-and-hold Sharpe, best ML Sharpe, best ML accuracy vs "
                         "the majority class. Change the instrument in the left sidebar (SPY → AAPL → MSFT) to show "
                         "that everything updates."),
    ("2 · Interactive backtest", "Select XGBoost. The green box confirms the settings match the report. Then: move "
                                 "the τ slider down (more days invested; at the bottom of the range several models "
                                 "simply stay invested every day, like buy-and-hold) and up (fewer days invested). "
                                 "Move the cost slider from 15 to 0 "
                                 "and to 40 bps and watch the Sharpe change. Switch on volatility-targeted sizing and "
                                 "show the max drawdown falling."),
    ("3 · Predictions", "Show the metrics table against the majority class, then the confusion matrix and the "
                        "calibration curve for one model."),
    ("4 · Walk-forward & robustness", "Show the 2016–2025 out-of-sample equity curves. The dashed buy-and-hold line "
                                      "ends on top."),
    ("5 · Train a model live", "Keep XGBoost and 'Use the hyperparameters tuned in the study' on, then press the "
                               "Download-and-train button (about 10 seconds). Say: 'this downloads today's data and "
                               "runs the whole pipeline; the numbers differ slightly from slide 16 because the data "
                               "now runs to the latest trading day.'"),
    ("6 · Figures & docs", "Browse a few generated figures to show that every chart in the report comes from code."),
    ("If something fails", "No internet: skip step 5 and show slide 16 instead. Dashboard not responding: refresh "
                           "the page; if needed restart it with the launch line below."),
]


# ---------------------------------------------------------------------------------------------


def shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para(doc, text="", *, size=10.5, bold=False, color=INK, italic=False, after=4, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    return p


def label(doc, text):
    p = para(doc, text.upper(), size=8.5, bold=True, color=AMBER, after=2)
    p.paragraph_format.space_before = Pt(6)
    return p


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)


def thumb(path: Path) -> io.BytesIO:
    im = Image.open(path).convert("RGB")
    im.thumbnail((1400, 788))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=85)
    buf.seek(0)
    return buf


def main() -> None:
    F = Facts()
    D = build(F)
    slides = list(D.prs.slides)
    info = content(F)
    assert len(info) == len(slides) == len(D.notes)
    pngs = [PREVIEW / f"slide-{i:02d}.png" for i in range(1, len(slides) + 1)]
    missing = [p.name for p in pngs if not p.exists()]
    if missing:
        raise SystemExit(f"run render_preview.py first; missing {missing[:3]}")

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, sz in (("Heading 1", 18), ("Heading 2", 14)):
        hs = doc.styles[name]
        hs.font.name, hs.font.size, hs.font.color.rgb = "Calibri", Pt(sz), INK

    # cover + how to use
    para(doc, "SPEAKER NOTES", size=10, bold=True, color=AMBER, after=2)
    para(doc, "Machine-Learning Trading Signals", size=24, bold=True, after=2)
    para(doc, "A comparative study against technical baseline strategies · mini-project viva", size=12,
         color=MUTED, after=10)
    para(doc, "Ishan Singh · Kenneth Martin   |   Mentors: Dr. Narendra V G · Dr. Ashwath Rao B", size=10.5,
         after=12)
    label(doc, "How to use this handbook")
    bullets(doc, [
        "One page per slide: what is on the slide, what to say, the numbers to remember and likely questions.",
        "'What to say' is the same text as the speaker notes inside the PowerPoint (Presenter View shows it).",
        "Every number here is generated from the project's results files, the same source as the slides.",
        "Appendix A: glossary · Appendix B: live-demo script · Appendix C: numbers cheat sheet.",
    ])
    words = [len(n.split()) for n in D.notes]
    total = sum(words) / WPM
    label(doc, f"Timing plan · about {total:.0f} minutes of talk at {WPM} words/minute, plus demo and questions")
    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(tbl.rows[0].cells, ["#", "Slide", "Time", "Clicks", "Presenter"]):
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.font.bold, r.font.size = True, Pt(9.5)
        shade(c, "F3E3C3")
    for i, (s, w, mo) in enumerate(zip(slides, words, D.motion.slides), 1):
        clicks = sum(1 for stp in mo["steps"] if stp["trigger"] == "click")
        row = tbl.add_row().cells
        vals = [str(i), s.shapes.title.text_frame.text.replace("\n", " ").replace("\v", " "),
                f"{w / WPM * 60:.0f} s", str(clicks) if clicks else "", ""]
        for c, v in zip(row, vals):
            c.text = ""
            c.paragraphs[0].add_run(v).font.size = Pt(9.5)
    widths = [Inches(0.35), Inches(4.3), Inches(0.6), Inches(0.6), Inches(0.95)]
    for row in tbl.rows:
        for c, wdt in zip(row.cells, widths):
            c.width = wdt
    para(doc, "Fill in the Presenter column to split the talk between the two of you.", size=9, color=MUTED,
         italic=True, after=0)

    # one page per slide
    for i, (s, info_i, note, mo, png) in enumerate(zip(slides, info, D.notes, D.motion.slides, pngs), 1):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        title = s.shapes.title.text_frame.text.replace("\n", " ").replace("\v", " ")
        doc.add_heading(f"Slide {i} · {title}", level=1)
        clicks = sum(1 for stp in mo["steps"] if stp["trigger"] == "click")
        trans = "Morph from the previous slide" if mo["transition"] == "morph" else "Fade"
        cue = f"~{len(note.split()) / WPM * 60:.0f} s · {trans}" + (f" · {clicks} clicks to build" if clicks else
                                                                       " · builds automatically")
        para(doc, cue, size=9.5, color=MUTED, italic=True, after=6)
        doc.add_picture(thumb(png), width=Inches(5.6))
        doc.paragraphs[-1].paragraph_format.space_after = Pt(6)
        label(doc, "What's on the slide")
        bullets(doc, info_i["shows"])
        label(doc, "What to say")
        para(doc, note.replace("[click]", "[CLICK]"), size=10.5, after=4)
        label(doc, "Numbers to remember")
        bullets(doc, info_i["numbers"])
        if info_i["qa"]:
            label(doc, "If they ask")
            for q, a in info_i["qa"]:
                p = doc.add_paragraph()
                r = p.add_run(f"Q: {q}  ")
                r.font.bold, r.font.size = True, Pt(10)
                r2 = p.add_run(a)
                r2.font.size = Pt(10)
                p.paragraph_format.space_after = Pt(4)

    # appendices
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Appendix A · Glossary", level=1)
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for term, d in GLOSSARY:
        c1, c2 = tbl.add_row().cells
        c1.text, c2.text = "", ""
        r = c1.paragraphs[0].add_run(term)
        r.font.bold, r.font.size = True, Pt(10)
        c2.paragraphs[0].add_run(d).font.size = Pt(10)
        c1.width, c2.width = Inches(1.7), Inches(5.0)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Appendix B · Live demo script (about 3–4 minutes)", level=1)
    para(doc, "Best placed at slide 15. Every number on screen is recomputed by the project's own code.",
         size=10, color=MUTED, italic=True)
    for step, text in DEMO:
        p = doc.add_paragraph()
        r = p.add_run(f"{step}. ")
        r.font.bold, r.font.size = True, Pt(10.5)
        p.add_run(text).font.size = Pt(10.5)
        p.paragraph_format.space_after = Pt(6)
    label(doc, "Launch (Windows, from the ML_Based_Algorithmic_Trading folder)")
    para(doc, r".venv\Scripts\python.exe -m streamlit run app\streamlit_app.py", size=10, after=2)
    para(doc, "Then open http://localhost:8501. First-time setup on a new machine (needs internet): "
              r"py -3.12 -m venv .venv   then   .venv\Scripts\python.exe -m pip install -r app\requirements.txt",
         size=9.5, color=MUTED, italic=True)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Appendix C · Numbers cheat sheet", level=1)
    sp = F.split
    lo_r, hi_r = F.vol_red
    est, lo, hi = F.bh_ci_spy
    rows = [
        ("Data", "SPY, AAPL, MSFT · Yahoo Finance daily OHLCV 2010–2025"),
        ("Splits", f"train {sp['train']['rows']:,} · validation {sp['val']['rows']} · test {sp['test']['rows']:,} "
                   "days (2-day purge)"),
        ("Features / target", f"{F.n_features} features in 5 groups · Down/Neutral/Up, ±{F.neutral * 100:.1f}% band, "
                              "next-open return"),
        ("Costs", f"{F.cost_total} bps per unit turnover ({F.fee_bps} fee + {F.slip_bps} slippage)"),
        ("Accuracy", f"models {F.acc_min * 100:.1f}–{F.acc_max * 100:.1f}% vs 'always Up' "
                     f"{F.maj_min * 100:.1f}–{F.maj_max * 100:.1f}% · only reliable difference SPY XGBoost (worse)"),
        ("Test Sharpe (B&H)", ", ".join(f"{t} {F.bh_sharpe[t]:.2f}" for t in TICKERS)
         + f" · SPY 95% CI {lo:.2f} to {hi:.2f}"),
        ("τ-rule vs B&H", f"{F.n_higher}/15 higher (none reliably) · {F.n_identical}/15 identical · 1 reliably worse"),
        ("Walk-forward", f"0/{F.wf_n} beat B&H · B&H Sharpe " + ", ".join(f"{t} {F.wf_table[t][0]:.2f}"
                                                                          for t in TICKERS)),
        ("Robustness", f"0 bps: {F.cost_beat_lo[0]}/15 above B&H · 40 bps: {F.cost_beat_hi[0]}/15 · vol targeting "
                       f"cuts max DD {lo_r * 100:.0f}–{hi_r * 100:.0f}%"),
        ("Engineering", f"{F.n_tests} tests · {F.minutes} min full run · {F.n_figures} figures · seed {F.seed}"),
        ("Live run (24 Oct)", f"{LIVE['seconds']} s · {LIVE['test_days']:,} test days · Sharpe {LIVE['sharpe']:.2f} "
                              f"vs {LIVE['bh_sharpe']:.2f}"),
    ]
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for k, v in rows:
        c1, c2 = tbl.add_row().cells
        c1.text, c2.text = "", ""
        r = c1.paragraphs[0].add_run(k)
        r.font.bold, r.font.size = True, Pt(10)
        shade(c1, "F3E3C3")
        c2.paragraphs[0].add_run(v).font.size = Pt(10)
        c1.width, c2.width = Inches(1.6), Inches(5.1)

    doc.core_properties.title = "Speaker notes: Machine-Learning Trading Signals"
    doc.core_properties.author = "Ishan Singh; Kenneth Martin"
    doc.core_properties.last_modified_by = "Ishan Singh; Kenneth Martin"
    doc.save(OUT)
    print(f"wrote {OUT}")

    if "--pdf" in sys.argv:
        import win32com.client as w

        word = w.gencache.EnsureDispatch("Word.Application")
        word.Visible = False
        d = word.Documents.Open(str(OUT), ReadOnly=True)
        try:
            d.ExportAsFixedFormat(str(OUT.with_suffix(".pdf")), 17)  # wdExportFormatPDF
        finally:
            d.Close(False)
        print(f"wrote {OUT.with_suffix('.pdf')}")


if __name__ == "__main__":
    main()
